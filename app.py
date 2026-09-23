import streamlit as st
import google.generativeai as genai
from pypdf import PdfReader
from PIL import Image
from docx import Document
import io
import time
import os

# ==========================================
# CẤU HÌNH GIAO DIỆN STREAMLIT
# ==========================================
st.set_page_config(page_title="PDF to Word OCR", page_icon="📚", layout="centered")

st.title("📚 Tool OCR Tiếng Việt (PDF to Word)")
st.caption("Sử dụng sức mạnh của Gemini AI - Đọc trực tiếp từ Secrets")

# ==========================================
# LẤY API KEYS TỪ STREAMLIT SECRETS & XOAY VÒNG
# ==========================================
class KeyRotator:
    def __init__(self, keys_text):
        raw_keys = keys_text.split('\n')
        self.keys = [k.strip() for k in raw_keys if k.strip()]
        self.index = 0
        
    def get_key(self):
        if not self.keys:
            return None
        key = self.keys[self.index]
        self.index = (self.index + 1) % len(self.keys)
        return key
    
    def count(self):
        return len(self.keys)

# Đọc chuỗi keys từ st.secrets (Key trong file secrets.toml đặt tên là GEMINI_KEYS)
secrets_keys_raw = st.secrets.get("GEMINI_KEYS", "")
rotator = KeyRotator(secrets_keys_raw)

# Sidebar kiểm tra trạng thái Key cho an tâm
with st.sidebar:
    st.header("⚙️ Trạng thái Hệ thống")
    if rotator.count() > 0:
        st.success(f"Đã nạp thành công {rotator.count()} API Keys từ Secrets.")
    else:
        st.error("Chưa tìm thấy `GEMINI_KEYS` trong Streamlit Secrets! Vui lòng kiểm tra lại cấu hình trên cloud.")

# ==========================================
# HÀM GỌI GEMINI API OCR TRANG ẢNH
# ==========================================
def process_page_ocr(img, rotator, max_retries=3):
    prompt = """Bạn là một chuyên gia OCR tiếng Việt. Hãy trích xuất chính xác toàn bộ văn bản từ hình ảnh này.
Yêu cầu bắt buộc:
1. Giữ nguyên từng từ, dấu câu của bản gốc.
2. Giữ nguyên việc ngắt dòng, phân đoạn (xuống dòng như trong ảnh).
3. Tuyệt đối KHÔNG thêm bất kỳ bình luận, giải thích, hay lời chào nào. Chỉ trả về đúng văn bản được trích xuất."""

    for attempt in range(max_retries):
        current_key = rotator.get_key()
        if not current_key:
            return "[Lỗi: Không tìm thấy API Key hợp lệ]"

        try:
            genai.configure(api_key=current_key)
            model = genai.GenerativeModel("gemini-1.5-flash") 
            response = model.generate_content([prompt, img])
            return response.text
            
        except Exception as e:
            time.sleep(2)
            if attempt == max_retries - 1:
                return f"\n[LỖI OCR TRANG NÀY: {str(e)}]\n"

# ==========================================
# HÀM TẠO FILE DOCX
# ==========================================
def create_docx(text_content):
    doc = Document()
    for paragraph in text_content.split('\n'):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
        else:
            doc.add_paragraph("")
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# ==========================================
# GIAO DIỆN CHÍNH
# ==========================================
if rotator.count() == 0:
    st.warning("⚠️ Vui lòng cấu hình `GEMINI_KEYS` trong mục Secrets của Streamlit Cloud trước khi sử dụng.")
else:
    uploaded_files = st.file_uploader("Kéo thả các file PDF vào đây (Mỗi part ~20 trang)", 
                                      type=["pdf"], 
                                      accept_multiple_files=True)

    if uploaded_files:
        st.write("---")
        
        for uploaded_file in uploaded_files:
            st.subheader(f"📄 Đang xử lý: `{uploaded_file.name}`")
            
            pdf_bytes = uploaded_file.getvalue()
            reader = PdfReader(io.BytesIO(pdf_bytes))
            total_pages = len(reader.pages)
            
            progress_bar = st.progress(0)
            status_text = st.empty()
            
            full_text = ""
            
            for page_num in range(total_pages):
                status_text.text(f"Đang OCR trang {page_num + 1} / {total_pages}...")
                
                page = reader.pages[page_num]
                images = page.images
                
                if images:
                    img_file_obj = images[0]
                    img = Image.open(io.BytesIO(img_file_obj.data))
                else:
                    img = Image.new('RGB', (800, 1000), color='white')
                
                page_text = process_page_ocr(img, rotator)
                full_text += page_text + "\n\n"
                
                progress_bar.progress((page_num + 1) / total_pages)
                
            status_text.text("✅ Đã OCR xong! Đang đóng gói file Word...")
            
            docx_bytes = create_docx(full_text)
            
            # Giữ nguyên tên file input đổi đuôi sang docx
            base_name = os.path.splitext(uploaded_file.name)[0]
            output_filename = f"{base_name}.docx"
            
            st.download_button(
                label=f"⬇️ Nhấn để tải về: {output_filename}",
                data=docx_bytes,
                file_name=output_filename,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=uploaded_file.name
            )
            st.write("---")
